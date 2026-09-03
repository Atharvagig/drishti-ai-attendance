import os
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MASTER_EXCEL_FILE = "Attendance_Master.xlsx"

def get_daily_excel_filename(date_str=None):
    """Generates standard daily Excel filename."""
    if not date_str:
        date_str = datetime.now().strftime('%Y-%m-%d')
    return f"Attendance_{date_str}.xlsx"

def sync_attendance_to_excel(student_id, name, department="General", date_str=None, time_str=None, confidence=0.0, email_sent=0):
    """
    Appends a new attendance record directly into both daily and master Excel spreadsheets with rich formatting.
    Columns: Date, Time, Student ID, Name, Department, Status, Confidence (%), Email Alert
    """
    if not date_str:
        date_str = datetime.now().strftime('%Y-%m-%d')
    if not time_str:
        time_str = datetime.now().strftime('%H:%M:%S')

    new_row = {
        'Date': str(date_str),
        'Time': str(time_str),
        'Student ID': str(student_id),
        'Name': str(name),
        'Department': str(department),
        'Status': 'Present',
        'Confidence (%)': f"{confidence:.1f}%" if isinstance(confidence, (int, float)) else str(confidence),
        'Email Alert': 'Sent' if email_sent else 'Skipped'
    }

    # Update Daily Excel File
    daily_filename = get_daily_excel_filename(date_str)
    _append_and_format_excel(daily_filename, new_row)

    # Update Master Excel File
    _append_and_format_excel(MASTER_EXCEL_FILE, new_row)

    print(f"[Excel Sync] Synchronized record for {name} ({student_id}) to {daily_filename} & {MASTER_EXCEL_FILE}")
    return True

def _append_and_format_excel(filepath, row_dict):
    """Helper method to append a row and apply openpyxl table styling."""
    columns = ['Date', 'Time', 'Student ID', 'Name', 'Department', 'Status', 'Confidence (%)', 'Email Alert']
    
    if os.path.exists(filepath):
        try:
            df = pd.read_excel(filepath)
        except Exception:
            df = pd.DataFrame(columns=columns)
    else:
        df = pd.DataFrame(columns=columns)

    # Prevent duplicate ID entries for the same date
    if not df.empty and ('Student ID' in df.columns) and ('Date' in df.columns):
        duplicate_check = df[(df['Student ID'].astype(str) == str(row_dict['Student ID'])) & 
                             (df['Date'].astype(str) == str(row_dict['Date']))]
        if not duplicate_check.empty:
            return  # Already exists in this Excel sheet

    new_entry_df = pd.DataFrame([row_dict])
    df = pd.concat([df, new_entry_df], ignore_index=True)
    df.to_excel(filepath, index=False)

    # Apply openpyxl professional formatting
    apply_openpyxl_styles(filepath)

def apply_openpyxl_styles(filepath):
    """Applies dark header theme, auto column width, status highlights and grid borders."""
    if not os.path.exists(filepath):
        return

    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    # Styling definitions
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid") # Dark Slate Gray
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    present_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid") # Light Emerald Green
    present_font = Font(name="Calibri", size=11, bold=True, color="065F46")
    
    regular_font = Font(name="Calibri", size=11, color="1F2937")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    # Apply Header Styles (Row 1)
    ws.row_dimensions[1].height = 26
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Apply Data Row Styles
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        ws.row_dimensions[row_idx].height = 22
        for cell_idx, cell in enumerate(row, start=1):
            cell.font = regular_font
            cell.border = thin_border
            
            # Alignments & Status Highlighting
            if cell_idx in [1, 2, 3, 7, 8]:  # Date, Time, ID, Confidence, Email
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            # Highlight 'Present' status in green badge style
            if cell.value == 'Present':
                cell.fill = present_fill
                cell.font = present_font
                cell.alignment = center_align

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(filepath)

if __name__ == '__main__':
    sync_attendance_to_excel('1001', 'Atharva Tripathi', 'Artificial Intelligence', confidence=99.2, email_sent=1)
    print("[Excel Styling] openpyxl formatting applied successfully.")
